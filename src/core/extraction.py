"""
File Extraction Module
----------------------
Provides a unified interface to extract raw text from multiple file
formats (docx, pptx, txt, pdf, images), yielding OCR results for PDFs and
images using specific parsers for Office documents and plain text files.
"""

import os
import mimetypes
from .logging_config import get_logger
from .ocr_client import OCRClient

logger = get_logger(__name__)

# Try imports
try:
    from pptx import Presentation

    PPTX_AVAILABLE = True
except ImportError:
    PPTX_AVAILABLE = False

try:
    from docx import Document

    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False


class FileExtractor:
    def __init__(self, ocr_client: OCRClient = None):
        self.ocr_client = ocr_client

    def extract(self, file_path: str, mime_type: str = None) -> str:
        """
        Extract raw text from file user specific handlers.
        """
        if not mime_type:
            mime_type, _ = mimetypes.guess_type(file_path)

        ext = os.path.splitext(file_path)[1].lower()

        # PDF / Images -> OCR
        if mime_type in [
            "application/pdf",
            "image/jpeg",
            "image/png",
            "image/bmp",
            "image/tiff",
        ] or ext in [".pdf", ".jpg", ".png"]:
            if not self.ocr_client:
                raise ValueError("OCR Client required for PDF/Image extraction")
            with open(file_path, "rb") as f:
                return self.ocr_client.extract_text(f)

        # TXT
        elif mime_type == "text/plain" or ext == ".txt":
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                return f.read()

        # PPTX
        elif (
            ext == ".pptx"
            or mime_type
            == "application/vnd.openxmlformats-officedocument.presentationml.presentation"
        ):
            if not PPTX_AVAILABLE:
                raise ImportError("python-pptx required for .pptx")
            return self._extract_pptx(file_path)

        # DOCX
        elif (
            ext == ".docx"
            or mime_type
            == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ):
            if not DOCX_AVAILABLE:
                raise ImportError("python-docx required for .docx")
            return self._extract_docx(file_path)

        else:
            raise ValueError(f"Unsupported file type: {mime_type} ({ext})")

    def _extract_pptx(self, file_path: str) -> str:
        prs = Presentation(file_path)
        text_content = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text_frame") and shape.text_frame:
                    text_content.append(shape.text_frame.text)
        return "\n".join(text_content)

    def _extract_docx(self, file_path: str) -> str:
        # Simplified extraction compared to original notebook specialized splitting
        # We just want raw text here. Detailed structure preservation might belong in chunking or specialized step.
        # But wait, the notebook had 'advanced processing' for headings.
        # Ideally, we return the text. If we want to preserve structure, we might need to return a structured object.
        # For now, let's extract all text cleanly.
        # If the 's01' contract is 'textual_data', maybe raw text is enough.

        with open(file_path, "rb") as f:
            doc = Document(f)

        full_text = []
        for para in doc.paragraphs:
            full_text.append(para.text)
        return "\n".join(full_text)
